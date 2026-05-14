"""
Convierte Informe_Ejecutivo.md a un documento Word (.docx) profesional.
"""
import docx
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from PIL import Image as PILImage
import io
import re
import os

def set_cell_shading(cell, color_hex):
    """Aplica color de fondo a una celda."""
    shading = docx.oxml.OxmlElement('w:shd')
    shading.set(qn('w:fill'), color_hex)
    shading.set(qn('w:val'), 'clear')
    cell._tc.get_or_add_tcPr().append(shading)

def style_table(table):
    """Aplica bordes y formato a una tabla."""
    tbl = table._tbl
    tblBorders = docx.oxml.OxmlElement('w:tblBorders')
    for border_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = docx.oxml.OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '4')
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), 'CCCCCC')
        tblBorders.append(border)
    tbl.tblPr.append(tblBorders)

def parse_md_to_docx(md_path, docx_path, portada_path=None):
    doc = docx.Document()
    
    # --- Configurar márgenes ---
    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    # --- Estilos base ---
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Calibri'
    style_normal.font.size = Pt(11)
    style_normal.font.color.rgb = RGBColor(0x1a, 0x1a, 0x2e)

    # --- PORTADA ---
    if portada_path and os.path.exists(portada_path):
        # Espacio superior
        for _ in range(3):
            doc.add_paragraph('')
        
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p_img.add_run()
        # Convertir a PNG en memoria si es webp
        try:
            img = PILImage.open(portada_path)
            buf = io.BytesIO()
            img.save(buf, format='PNG')
            buf.seek(0)
            run.add_picture(buf, width=Inches(4))
        except Exception as e:
            print(f"⚠ No se pudo insertar la portada: {e}")
        
        doc.add_paragraph('')

    # Título de portada
    for _ in range(2 if portada_path else 6):
        doc.add_paragraph('')
    
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p_title.add_run('INFORME EJECUTIVO DE AVANCE DEL PROYECTO')
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor(0x0d, 0x94, 0x67)
    
    doc.add_paragraph('')
    
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p_sub.add_run('YALOTENGO')
    run.bold = True
    run.font.size = Pt(28)
    run.font.color.rgb = RGBColor(0x1a, 0x1a, 0x2e)
    
    p_sub2 = doc.add_paragraph()
    p_sub2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p_sub2.add_run('Ecosistema Web de Comercialización Tecnológica')
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0x6b, 0x72, 0x80)
    
    doc.add_paragraph('')
    doc.add_paragraph('')
    
    p_empresa = doc.add_paragraph()
    p_empresa.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p_empresa.add_run('Empresa: Invéntalo')
    run.font.size = Pt(13)
    run.bold = True
    
    p_fecha = doc.add_paragraph()
    p_fecha.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p_fecha.add_run('Fecha: 13 de mayo de 2026')
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x6b, 0x72, 0x80)
    
    p_metodo = doc.add_paragraph()
    p_metodo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p_metodo.add_run('Metodología: Scrum (Desarrollo Ágil)')
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x6b, 0x72, 0x80)
    
    doc.add_page_break()

    # --- Leer el archivo MD ---
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()
    
    i = 0
    while i < len(lines):
        line = lines[i].rstrip('\r\n')
        
        # Saltar título principal y datos generales de portada (ya los pusimos)
        if line.startswith('# INFORME EJECUTIVO'):
            i += 1
            continue
        if line.startswith('**Proyecto:**') or line.startswith('**Empresa:**') or line.startswith('**Fecha:**'):
            i += 1
            continue
        
        # Líneas separadoras
        if line.strip() == '---':
            i += 1
            continue
        
        # Líneas vacías
        if not line.strip():
            i += 1
            continue
        
        # Encabezados
        if line.startswith('## '):
            heading_text = line[3:].strip()
            # Limpiar negritas en headings
            heading_text = heading_text.replace('**', '')
            h = doc.add_heading(heading_text, level=1)
            h.runs[0].font.color.rgb = RGBColor(0x0d, 0x94, 0x67)
            i += 1
            continue
        
        if line.startswith('### '):
            heading_text = line[4:].strip()
            heading_text = heading_text.replace('**', '')
            doc.add_heading(heading_text, level=2)
            i += 1
            continue
        
        if line.startswith('#### '):
            heading_text = line[5:].strip()
            heading_text = heading_text.replace('**', '')
            doc.add_heading(heading_text, level=3)
            i += 1
            continue
        
        # Tablas
        if '|' in line and line.strip().startswith('|'):
            table_lines = []
            while i < len(lines) and '|' in lines[i] and lines[i].strip().startswith('|'):
                table_lines.append(lines[i].rstrip('\r\n'))
                i += 1
            
            # Parsear tabla
            if len(table_lines) >= 2:
                # Extraer filas (saltar separador ---)
                rows_data = []
                for tl in table_lines:
                    cells = [c.strip() for c in tl.split('|')[1:-1]]
                    if cells and not all(re.match(r'^:?-+:?$', c) for c in cells):
                        rows_data.append(cells)
                
                if rows_data:
                    num_cols = len(rows_data[0])
                    table = doc.add_table(rows=len(rows_data), cols=num_cols)
                    table.alignment = WD_TABLE_ALIGNMENT.CENTER
                    style_table(table)
                    
                    for row_idx, row_data in enumerate(rows_data):
                        for col_idx, cell_text in enumerate(row_data):
                            if col_idx < num_cols:
                                cell = table.cell(row_idx, col_idx)
                                # Limpiar markdown bold
                                clean_text = cell_text.replace('**', '').strip()
                                cell.text = clean_text
                                
                                # Header row styling
                                if row_idx == 0:
                                    set_cell_shading(cell, '0d9467')
                                    for paragraph in cell.paragraphs:
                                        for run in paragraph.runs:
                                            run.bold = True
                                            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                                            run.font.size = Pt(10)
                                elif row_idx % 2 == 0:
                                    set_cell_shading(cell, 'F3F4F6')
                    
                    doc.add_paragraph('')  # Espacio después de tabla
            continue
        
        # Listas con viñetas
        if line.startswith('- '):
            text = line[2:].strip()
            # Checkbox items
            text = text.replace('[ ] ', '☐ ').replace('[x] ', '☑ ')
            # Limpiar negritas
            p = doc.add_paragraph(style='List Bullet')
            # Parsear negritas inline
            parts = re.split(r'(\*\*.*?\*\*)', text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                else:
                    p.add_run(part)
            i += 1
            continue
        
        # Texto con formato itálico al final
        if line.startswith('*') and line.endswith('*'):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(line.strip('*'))
            run.italic = True
            run.font.color.rgb = RGBColor(0x6b, 0x72, 0x80)
            i += 1
            continue
        
        # Párrafo normal
        p = doc.add_paragraph()
        # Parsear negritas inline
        parts = re.split(r'(\*\*.*?\*\*)', line)
        for part in parts:
            if part.startswith('**') and part.endswith('**'):
                run = p.add_run(part[2:-2])
                run.bold = True
            else:
                p.add_run(part)
        i += 1
    
    # --- Guardar ---
    doc.save(docx_path)
    print(f"[OK] Documento generado: {docx_path}")

if __name__ == '__main__':
    base = r"c:\Users\Danilo\Documents\Invéntalo\yalotengo-master"
    md_file = os.path.join(base, "Documentación", "Informe_Ejecutivo.md")
    docx_file = os.path.join(base, "Documentación", "Informe_Ejecutivo_Yalotengo.docx")
    portada = os.path.join(base, "frontend", "public", "portada.webp")
    
    # Si no se puede usar webp, intentar jpg
    if not os.path.exists(portada):
        portada = os.path.join(base, "frontend", "public", "portada.jpg")
    
    parse_md_to_docx(md_file, docx_file, portada_path=portada)
