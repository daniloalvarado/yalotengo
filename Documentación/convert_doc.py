import docx
import sys
import os

def convert_docx_to_md(docx_path, md_path):
    doc = docx.Document(docx_path)
    with open(md_path, 'w', encoding='utf-8') as f:
        for para in doc.paragraphs:
            # Basic conversion: heading levels
            if para.style.name.startswith('Heading'):
                level = para.style.name.split(' ')[-1]
                try:
                    level = int(level)
                    f.write('#' * level + ' ' + para.text + '\n\n')
                except:
                    f.write('# ' + para.text + '\n\n')
            elif para.text.strip():
                f.write(para.text + '\n\n')
        
        # Also try to extract tables
        for table in doc.tables:
            f.write('\n| ' + ' | '.join([cell.text.replace('\n', ' ') for cell in table.rows[0].cells]) + ' |\n')
            f.write('| ' + ' | '.join(['---' for _ in table.rows[0].cells]) + ' |\n')
            for row in table.rows[1:]:
                f.write('| ' + ' | '.join([cell.text.replace('\n', ' ') for cell in row.cells]) + ' |\n')
            f.write('\n')

if __name__ == "__main__":
    docx_file = r"c:\Users\Danilo\Documents\Invéntalo\yalotengo-master\Documentación\Documentación Yalotengo.docx"
    md_file = r"c:\Users\Danilo\Documents\Invéntalo\yalotengo-master\Documentación\Documentación Yalotengo.md"
    convert_docx_to_md(docx_file, md_file)
    print(f"Converted {docx_file} to {md_file}")
